import tkinter as tk
from tkinter import ttk, messagebox
from datetime import datetime, timedelta
import configparser
from tkcalendar import DateEntry
from .utils import format_currency, show_error_message
from .company.reporting import ReportingHandler

class CompanyModule:
    def __init__(self, parent, setup_ui=True):
        self.parent = parent
        
        # Read configuration
        self.config = configparser.ConfigParser()
        self.config.read('Config/config.ini')
        self.debug_mode = self.config.getboolean('DEBUG', 'debugmode', fallback=False)
        
        self.reporting_handler = ReportingHandler(debug_mode=self.debug_mode)
        
        if setup_ui:
            self.setup_ui()
    
    def setup_ui(self):
        # Main frame
        main_frame = ttk.Frame(self.parent, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # Report parameters
        params_frame = ttk.LabelFrame(main_frame, text="Report Parameters", padding="10")
        params_frame.pack(fill=tk.X, pady=(0, 10))
        
        # Date range
        ttk.Label(params_frame, text="From Date:").grid(row=0, column=0, sticky=tk.W, pady=2)
        self.from_date_entry = DateEntry(params_frame, width=12, background='darkblue', foreground='white', borderwidth=2, date_pattern='y-mm-dd')
        self.from_date_entry.set_date(datetime.now() - timedelta(days=30))
        self.from_date_entry.grid(row=0, column=1, sticky=(tk.W, tk.E), pady=2, padx=(5, 10))
        
        ttk.Label(params_frame, text="To Date:").grid(row=0, column=2, sticky=tk.W, pady=2)
        self.to_date_entry = DateEntry(params_frame, width=12, background='darkblue', foreground='white', borderwidth=2, date_pattern='y-mm-dd')
        self.to_date_entry.set_date(datetime.now())
        self.to_date_entry.grid(row=0, column=3, sticky=(tk.W, tk.E), pady=2, padx=(5, 10))
        
        # Report type
        ttk.Label(params_frame, text="Report Type:").grid(row=1, column=0, sticky=tk.W, pady=(10, 2))
        self.report_type_var = tk.StringVar()
        report_type_combo = ttk.Combobox(params_frame, textvariable=self.report_type_var, state="readonly")
        report_type_combo['values'] = ["Daily", "Weekly", "Monthly", "Custom Range"]
        report_type_combo.grid(row=1, column=1, sticky=(tk.W, tk.E), pady=(10, 2), padx=(5, 10))
        report_type_combo.set("Monthly")
        
        # Generate button
        ttk.Button(params_frame, text="Generate Report", command=self.generate_report).grid(row=1, column=3, sticky=tk.E, pady=(10, 2))
        
        # Results frame
        results_frame = ttk.LabelFrame(main_frame, text="Report Results", padding="10")
        results_frame.pack(fill=tk.BOTH, expand=True)
        
        # Create text widget with scrollbar
        text_frame = ttk.Frame(results_frame)
        text_frame.pack(fill=tk.BOTH, expand=True)
        
        self.results_text = tk.Text(text_frame, wrap=tk.WORD)
        scrollbar = ttk.Scrollbar(text_frame, orient=tk.VERTICAL, command=self.results_text.yview)
        self.results_text.configure(yscrollcommand=scrollbar.set)
        
        self.results_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # Export button
        ttk.Button(results_frame, text="Export Report", command=self.export_report).pack(pady=(10, 0))
        
        # Configure grid weights
        params_frame.columnconfigure(1, weight=1)
        params_frame.columnconfigure(3, weight=1)
        main_frame.columnconfigure(0, weight=1)
        main_frame.rowconfigure(1, weight=1)
    
    def generate_report(self):
        """Generate company report"""
        if self.debug_mode:
            print("--- Generating Company Report (Debug Mode) ---")
        
        try:
            from_date = self.from_date_entry.get_date().strftime("%Y-%m-%d")
            to_date = self.to_date_entry.get_date().strftime("%Y-%m-%d")
            report_type = self.report_type_var.get()
            
            if self.debug_mode:
                print(f"Report parameters: from={from_date}, to={to_date}, type={report_type}")

            # Validate dates
            if self.from_date_entry.get_date() > self.to_date_entry.get_date():
                show_error_message("Error", "From date cannot be after to date")
                return
        except ValueError:
            show_error_message("Error", "Please enter valid dates (YYYY-MM-DD)")
            print("Error: Please enter valid dates (YYYY-MM-DD)")
            return
        
        # Clear previous results
        self.results_text.delete(1.0, tk.END)
        
        # Generate report header
        report_header = f"""
ICU MANAGEMENT SYSTEM - COMPANY REPORT
Report Period: {from_date} to {to_date}
Report Type: {report_type}
Generated on: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}

""" + "="*80 + "\n\n"
        
        self.results_text.insert(tk.END, report_header)
        
        # Get patient revenues and pass-through costs
        if self.debug_mode:
            print("Calculating patient revenues...")
        patient_revenues = self.reporting_handler.calculate_patient_revenues(from_date, to_date)
        total_patient_revenue = patient_revenues['total']
        pass_through_costs = patient_revenues['operational_cost']
        if self.debug_mode:
            print(f"Patient revenues calculated: Total={total_patient_revenue}, Pass-through={pass_through_costs}")
            print("Patient details:", patient_revenues['details'])

        # Get staff costs
        if self.debug_mode:
            print("Calculating doctor costs...")
        doctor_costs = self.reporting_handler.calculate_doctor_costs(from_date, to_date)
        if self.debug_mode:
            print(f"Doctor costs calculated: Total={doctor_costs['total']}")
            print("Doctor details:", doctor_costs['details'])
            
        if self.debug_mode:
            print("Calculating nurse costs...")
        nurse_costs = self.reporting_handler.calculate_nurse_costs(from_date, to_date)
        if self.debug_mode:
            print(f"Nurse costs calculated: Total={nurse_costs['total']}")
            print("Nurse details:", nurse_costs['details'])
            
        total_staff_cost = doctor_costs['total'] + nurse_costs['total']
        if self.debug_mode:
            print(f"Total staff cost: {total_staff_cost}")

        # Calculate total operational costs
        total_operational_cost = total_staff_cost + pass_through_costs
        if self.debug_mode:
            print(f"Total operational cost: {total_operational_cost}")
        
        # Calculate net profit
        net_profit = total_patient_revenue - total_operational_cost
        if self.debug_mode:
            print(f"Net profit: {net_profit}")
            print("--- Report Generation Complete ---")

        # Display summary
        summary_text = f"""
FINANCIAL SUMMARY
=================
Total Patient Revenue: {format_currency(total_patient_revenue)}

Operational Costs:
  - Staff Costs: {format_currency(total_staff_cost)}
  - Pass-Through Costs (Labs, Drugs, etc.): {format_currency(pass_through_costs)}
-----------------
Total Operational Costs: {format_currency(total_operational_cost)}

Net Profit/Loss: {format_currency(net_profit)}

"""
        
        self.results_text.insert(tk.END, summary_text)
        
        # Display doctor details
        self.results_text.insert(tk.END, "\nDOCTORS\n" + "-"*40 + "\n")
        for doctor in doctor_costs['details']:
            self.results_text.insert(tk.END, f"{doctor['name']}: {format_currency(doctor['cost'])}\n")
        
        # Display nurse details
        self.results_text.insert(tk.END, "\nNURSES\n" + "-"*40 + "\n")
        for nurse in nurse_costs['details']:
            self.results_text.insert(tk.END, f"{nurse['name']} ({nurse['level']}): {format_currency(nurse['cost'])}\n")
        
        # Display patient details
        self.results_text.insert(tk.END, "\nPATIENTS\n" + "-"*40 + "\n")
        for patient in patient_revenues['details']:
            self.results_text.insert(tk.END, f"{patient['name']}: {format_currency(patient['revenue'])}\n")
    
    
    def export_report(self):
        """Export report to file"""
        report_content = self.results_text.get(1.0, tk.END)
        if not report_content.strip():
            messagebox.showwarning("Warning", "Please generate a report first")
            return

        import openpyxl
        from datetime import datetime as dt
        filename = f"company_report_{dt.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        try:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Company Report"

            for line in report_content.split('\n'):
                sheet.append([line])

            workbook.save(filename)
            messagebox.showinfo("Export Success", f"Report exported to {filename}")
        except Exception as e:
            show_error_message("Export Error", f"Failed to export report: {e}")
            print(f"Export Error: Failed to export report: {e}")
